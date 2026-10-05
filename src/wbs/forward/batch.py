"""Structure-tree batch ([D2] 正向构造): sample → scene.json (pre-render dynamic gate, reseed) → manifest."""

from __future__ import annotations

import csv
import dataclasses
import time
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from ..config import Settings
from ..jsonio import read_json, write_json, write_text
from ..layout import JobPaths, Workspace
from ..prompts import render
from ..providers import CallContext, Providers
from ..providers.mock import register_mock
from ..qc.pregate import pregate, write_pregate
from ..taxonomy import ControlSpec, sample_controls
from . import texts
from .procedural import generate

MANIFEST_COLUMNS = [
    ("job_id", "任务ID"), ("title", "标题"), ("route", "路线"), ("shot_form", "镜头形式"), ("shots", "镜头数"),
    ("subject", "主体"), ("color", "颜色"), ("era", "时代"), ("viewpoint", "视角"), ("camera_move", "运镜"),
    ("content_class", "内容类别"), ("seed", "种子"), ("pregate", "渲染前动态闸门"), ("qc", "质检结论"),
    ("video", "视频文件"), ("status", "状态"),
]


def plan_job(ws: Workspace, settings: Settings, batch_id: str, control: ControlSpec,
             spec_name: str = "forward_batch", max_attempts: int = 5) -> JobPaths:
    spec = settings.spec(spec_name)
    job_id = f"{batch_id}_{control.index:04d}"
    job = ws.job(batch_id, job_id)
    job.root.mkdir(parents=True, exist_ok=True)
    attempts: list[dict[str, Any]] = []
    current = control
    for k in range(max_attempts):
        scene = generate(current, spec, job_id).to_json_dict()
        checked = pregate(scene, settings)
        gate = checked["gate"]
        attempts.append({"seed": current.seed, "dynamic_gate": gate["status"], "issue_counts": gate["issue_counts"],
                         "framing_precheck": checked["framing"]["status"], "grammar": checked["grammar"]["status"]})
        if checked["status"] == "passed":
            break
        current = dataclasses.replace(control, seed=(control.seed + 7919 * (k + 1)) % 2**31)
    checked["gate"] = {**gate, "status": checked["status"]}
    write_json(job.scene, scene)
    write_json(job.control_spec, {**control.as_dict(), "labels": control.labels(), "final_seed": current.seed})
    write_pregate(job, checked, attempts=attempts)
    control_text, control_ref = texts.control_layer_prompt(scene)
    v2v_text, v2v_ref = texts.v2v_render_prompt(scene)
    write_text(job.prompts_dir / "白模控制层.txt", control_text)
    write_text(job.prompts_dir / "V2V渲染提示词.txt", v2v_text)
    job.update_meta(job_id=job_id, batch_id=batch_id, route="forward_batch", spec=spec_name, title=scene["title"],
                    control=control.as_dict(), seed=current.seed, shots=len(scene["shots"]),
                    pregate=checked["status"], grammar=checked["grammar"]["status"], pregate_attempts=len(attempts),
                    status="planned",
                    prompt_refs={"control_layer": control_ref, "v2v_render": v2v_ref})
    return job


REWRITE_SYSTEM = "你是 3D 预演建模指令编写助手，只输出改写后的提示词正文。"


def rewrite_prompts(jobs: list[JobPaths], providers: Providers) -> dict[str, Any]:
    """[D2] 控制层 prompt → 具体建模提示词（LLM 改写；默认 mock）。结果只作为人工/代理建模的参考文本。"""
    done = 0
    for job in jobs:
        control = (job.prompts_dir / "白模控制层.txt").read_text(encoding="utf-8")
        user, ref = render("forward.control_rewrite", control_prompt=control)
        result = providers.llm.complete(CallContext(job.key, "control_rewrite", job.batch_id), system=REWRITE_SYSTEM,
                                        user=user, task="rewrite_control_prompt", payload={"control_prompt": control})
        write_text(job.prompts_dir / "建模提示词.txt", result.text.strip() + "\n")
        job.update_meta(rewrite={"prompt_ref": ref, "model": result.model, "simulated": result.simulated})
        done += 1
    return {"rewritten": done}


@register_mock("rewrite_control_prompt")
def _mock_rewrite(payload: dict[str, Any]) -> str:
    control = str(payload.get("control_prompt", ""))
    body = control.replace("目标：\n", "").strip()
    return ("（mock 改写，未调用真实模型；数字、结构树控制信息与白模规则原样保留）\n"
            "请在 Blender 中按以下要求直接搭建并渲染低精度 blocking 预演：先建地面与大结构体块，再建主体几何代理并打整体位移与"
            "转向关键帧，最后按镜头表建相机与切换标记，不添加材质、贴图与特效。\n\n" + body)


def plan_batch(ws: Workspace, settings: Settings, batch_id: str, n: int, seed: int, narrative_ratio: float = 0.5,
               subjects: list[str] | None = None, spec_name: str = "forward_batch") -> list[JobPaths]:
    controls = sample_controls(n, seed, narrative_ratio=narrative_ratio, subjects=subjects)
    batch_dir = ws.batch_dir(batch_id)
    batch_dir.mkdir(parents=True, exist_ok=True)
    write_json(batch_dir / "batch.json", {"batch_id": batch_id, "route": "forward_batch", "n": n, "seed": seed,
                                          "narrative_ratio": narrative_ratio, "subjects": subjects, "spec": spec_name,
                                          "created_at": time.strftime("%Y-%m-%dT%H:%M:%S")})
    jobs = [plan_job(ws, settings, batch_id, c, spec_name) for c in controls]
    write_manifest(batch_dir, ws.iter_jobs(batch_id))
    return jobs


def manifest_rows(jobs: list[JobPaths]) -> list[dict[str, Any]]:
    rows = []
    for job in jobs:
        meta = job.read_meta()
        control = meta.get("control") or {}
        labels = {}
        if job.control_spec.exists():
            labels = read_json(job.control_spec).get("labels", {})
        rows.append({
            "job_id": job.job_id, "title": meta.get("title", ""), "route": meta.get("route", ""),
            "shot_form": labels.get("shot_form_label", ""), "shots": meta.get("shots", ""),
            "subject": labels.get("subject_label", meta.get("subject", "")), "color": labels.get("color_label", ""),
            "era": labels.get("era_label", ""), "viewpoint": labels.get("viewpoint_label", ""),
            "camera_move": labels.get("camera_move_label", ""),
            "content_class": labels.get("content_class_label", meta.get("content_class", "")),
            "seed": meta.get("seed", control.get("seed", "")), "pregate": meta.get("pregate", ""),
            "qc": meta.get("qc_status", "not_run"), "video": meta.get("video", ""), "status": meta.get("status", ""),
        })
    return rows


def write_manifest(batch_dir: Path, jobs: list[JobPaths]) -> tuple[Path, Path]:
    rows = manifest_rows(jobs)
    csv_path, xlsx_path = batch_dir / "批次清单.csv", batch_dir / "批次清单.xlsx"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([label for _, label in MANIFEST_COLUMNS])
        for row in rows:
            writer.writerow([row[key] for key, _ in MANIFEST_COLUMNS])
    wb = Workbook()
    ws = wb.active
    ws.title = "批次清单"
    ws.append([label for _, label in MANIFEST_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row[key] for key, _ in MANIFEST_COLUMNS])
    for i, (key, label) in enumerate(MANIFEST_COLUMNS, start=1):
        width = max([len(str(label))] + [len(str(r[key])) for r in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width * 1.6, 8), 60)
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return csv_path, xlsx_path
